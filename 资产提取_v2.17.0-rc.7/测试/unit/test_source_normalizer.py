#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""工单01-Task2：统一来源数据模型（规格 4）。

统一中间文件为 UTF-8 Markdown，正文前带结构化来源清单；每个规范化段落可追溯到
原文件、文件哈希、卷/章/集/场、段落序号和原文摘句。
"""
import sys
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(REPO_ROOT))

from scripts.engine.source_normalizer import (  # noqa: E402
    NormalizedSource,
    SourceNormalizeError,
    SourceUnit,
    normalize_sources,
)


def test_txt_and_markdown_keep_order_hash_and_locator(tmp_path):
    a = tmp_path / "卷1.txt"
    a.write_text("第一章\n甲出现。", encoding="utf-8")
    b = tmp_path / "卷2.md"
    b.write_text("# 第二章\n乙出现。", encoding="utf-8")
    result = normalize_sources([a, b])
    assert [u.order for u in result.units] == [1, 2]
    assert result.units[0].locator == "文件1/段落1"
    assert len(result.units[0].source_sha256) == 64


def test_normalized_returns_declared_types(tmp_path):
    a = tmp_path / "a.txt"
    a.write_text("正文甲", encoding="utf-8")
    result = normalize_sources([a])
    assert isinstance(result, NormalizedSource)
    assert all(isinstance(u, SourceUnit) for u in result.units)
    assert isinstance(result.manifest, dict)
    assert result.markdown.endswith("\n")


def test_source_comment_is_written_before_each_source(tmp_path):
    a = tmp_path / "卷1.txt"
    a.write_text("第一章\n甲出现。", encoding="utf-8")
    b = tmp_path / "卷2.md"
    b.write_text("第二章\n乙出现。", encoding="utf-8")
    result = normalize_sources([a, b])
    digest_a = result.units[0].source_sha256
    assert f"<!-- source: 卷1.txt sha256:{digest_a} order:1 -->" in result.markdown
    assert result.markdown.count("<!-- source:") == 2


def test_quote_can_be_traced_back_into_normalized_markdown(tmp_path):
    a = tmp_path / "卷1.txt"
    a.write_text("第一章\n甲在茶楼遇见乙。\n\n他递出一枚玉牌。", encoding="utf-8")
    result = normalize_sources([a])
    for unit in result.units:
        assert unit.text.strip()
        assert unit.text.strip() in result.markdown
    assert "他递出一枚玉牌。" in result.markdown


def test_paragraph_numbering_is_one_based_and_monotonic(tmp_path):
    a = tmp_path / "卷1.txt"
    a.write_text("第一段\n\n第二段\n\n第三段", encoding="utf-8")
    result = normalize_sources([a])
    assert [u.locator for u in result.units] == ["文件1/段落1", "文件1/段落2", "文件1/段落3"]


def test_source_sha256_matches_file_bytes(tmp_path):
    import hashlib

    a = tmp_path / "卷1.txt"
    a.write_bytes("第一章\n甲出现。".encode("utf-8"))
    result = normalize_sources([a])
    assert result.units[0].source_sha256 == hashlib.sha256(a.read_bytes()).hexdigest()
    assert result.manifest["sources"][0]["sha256"] == result.units[0].source_sha256
    assert result.manifest["sources"][0]["bytes"] == a.stat().st_size


def test_natural_sort_used_when_no_order_file(tmp_path):
    names = ["卷10.txt", "卷2.txt", "卷1.txt"]
    paths = []
    for n in names:
        p = tmp_path / n
        p.write_text(f"{n} 内容", encoding="utf-8")
        paths.append(p)
    result = normalize_sources(paths)
    assert [Path(m["source_path"]).name for m in result.manifest["sources"]] == [
        "卷1.txt", "卷2.txt", "卷10.txt"
    ]
    assert result.manifest["order_source"] == "natural"


def test_explicit_order_file_wins(tmp_path):
    for n in ("卷1.txt", "卷2.txt", "卷3.txt"):
        (tmp_path / n).write_text(f"{n} 内容", encoding="utf-8")
    order = tmp_path / "顺序清单.txt"
    order.write_text("卷3.txt\n卷1.txt\n卷2.txt\n", encoding="utf-8")
    result = normalize_sources([tmp_path / "卷1.txt", tmp_path / "卷2.txt", tmp_path / "卷3.txt"],
                               order_file=order)
    assert [Path(m["source_path"]).name for m in result.manifest["sources"]] == [
        "卷3.txt", "卷1.txt", "卷2.txt"
    ]
    assert result.manifest["order_source"] == "explicit"
    assert result.units[0].locator == "文件1/段落1"


def test_order_file_duplicate_rejected(tmp_path):
    (tmp_path / "卷1.txt").write_text("甲", encoding="utf-8")
    order = tmp_path / "顺序清单.txt"
    order.write_text("卷1.txt\n卷1.txt\n", encoding="utf-8")
    with pytest.raises(SourceNormalizeError) as ctx:
        normalize_sources([tmp_path / "卷1.txt"], order_file=order)
    assert ctx.value.code == "SOURCE_ORDER_DUPLICATE"


def test_order_file_missing_source_rejected(tmp_path):
    (tmp_path / "卷1.txt").write_text("甲", encoding="utf-8")
    (tmp_path / "卷2.txt").write_text("乙", encoding="utf-8")
    order = tmp_path / "顺序清单.txt"
    order.write_text("卷1.txt\n卷3.txt\n", encoding="utf-8")
    with pytest.raises(SourceNormalizeError) as ctx:
        normalize_sources([tmp_path / "卷1.txt", tmp_path / "卷2.txt"], order_file=order)
    assert ctx.value.code == "SOURCE_ORDER_MISSING"


def test_order_file_outside_root_rejected(tmp_path):
    root = tmp_path / "项目"
    root.mkdir()
    src = root / "卷1.txt"
    src.write_text("甲", encoding="utf-8")
    outside = tmp_path / "外部.txt"
    outside.write_text("乙", encoding="utf-8")
    order = root / "顺序清单.txt"
    order.write_text("../外部.txt\n", encoding="utf-8")
    with pytest.raises(SourceNormalizeError) as ctx:
        normalize_sources([src, outside], order_file=order)
    assert ctx.value.code == "SOURCE_ORDER_OUTSIDE_ROOT"


def test_gb18030_is_detected_explicitly(tmp_path):
    a = tmp_path / "卷1.txt"
    a.write_bytes("第一章\n甲出现。".encode("gb18030"))
    result = normalize_sources([a])
    assert "甲出现。" in result.markdown
    assert result.manifest["sources"][0]["encoding"] == "gb18030"


def test_utf8_bom_is_stripped(tmp_path):
    a = tmp_path / "卷1.txt"
    a.write_bytes("\ufeff第一章\n甲出现。".encode("utf-8"))
    result = normalize_sources([a])
    assert result.markdown.lstrip().startswith("<!-- source:")
    assert "\ufeff" not in result.markdown
    assert "第一章" in result.markdown


def test_undecodable_bytes_fail_loudly(tmp_path):
    a = tmp_path / "卷1.txt"
    a.write_bytes(b"\xff\xfe\x00\x01\xff\xfd")
    with pytest.raises(SourceNormalizeError) as ctx:
        normalize_sources([a])
    assert ctx.value.code == "SOURCE_DECODE_FAILED"


def test_pdf_rejected_before_reading_content(tmp_path):
    a = tmp_path / "剧本.pdf"
    a.write_bytes(b"%PDF-1.7\nnot really a pdf")
    with pytest.raises(SourceNormalizeError) as ctx:
        normalize_sources([a])
    assert ctx.value.code == "UNSUPPORTED_SOURCE_FORMAT"
    assert "TXT" in ctx.value.detail or "txt" in ctx.value.detail


def test_unknown_suffix_rejected(tmp_path):
    a = tmp_path / "剧本.rtf"
    a.write_text("内容", encoding="utf-8")
    with pytest.raises(SourceNormalizeError) as ctx:
        normalize_sources([a])
    assert ctx.value.code == "UNSUPPORTED_SOURCE_FORMAT"


def test_missing_file_rejected(tmp_path):
    with pytest.raises(SourceNormalizeError) as ctx:
        normalize_sources([tmp_path / "不存在.txt"])
    assert ctx.value.code == "SOURCE_NOT_FOUND"


def test_empty_source_list_rejected():
    with pytest.raises(SourceNormalizeError) as ctx:
        normalize_sources([])
    assert ctx.value.code == "SOURCE_LIST_EMPTY"


def test_markdown_output_is_deterministic(tmp_path):
    a = tmp_path / "卷1.txt"
    a.write_text("第一章\n甲出现。", encoding="utf-8")
    first = normalize_sources([a])
    second = normalize_sources([a])
    assert first.markdown == second.markdown
    assert first.manifest == second.manifest


def test_paths_are_normalized_to_forward_slashes(tmp_path):
    a = tmp_path / "卷1.txt"
    a.write_text("甲", encoding="utf-8")
    result = normalize_sources([a])
    assert "\\" not in result.manifest["sources"][0]["source_path"]


# ---------------------------------------------------------------------------
# Task 3：DOCX / EPUB 正式输入，PDF 明确拒绝（规格 2.2 / 4）
# ---------------------------------------------------------------------------

FIXTURES = REPO_ROOT / "测试" / "fixtures" / "sources"
DOCX_FIXTURE = FIXTURES / "最小剧本.docx"
EPUB_FIXTURE = FIXTURES / "最小小说.epub"


def test_docx_fixture_exists_and_is_structured_as_expected():
    """夹具自身可核验：python-docx 读出的标题/正文/表格顺序即规范化期望顺序。"""
    from docx import Document

    assert DOCX_FIXTURE.is_file(), f"缺少夹具 {DOCX_FIXTURE}"
    doc = Document(str(DOCX_FIXTURE))
    texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    assert texts[0] == "第一场 茶楼"
    assert "甲走进茶楼，选了个靠窗的座位。" in texts
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[1].cells[0].text.strip() == "玉牌"


def test_epub_fixture_exists_and_is_structured_as_expected():
    from ebooklib import epub

    assert EPUB_FIXTURE.is_file(), f"缺少夹具 {EPUB_FIXTURE}"
    book = epub.read_epub(str(EPUB_FIXTURE), options={"ignore_ncx": True})
    names = []
    for entry in book.spine:
        idref = entry[0] if isinstance(entry, tuple) else entry
        item = book.get_item_with_id(idref)
        if item is not None:
            names.append((item.get_name(), item.is_chapter()))
    # spine 顺序与文件名字母序相反（chapter_b 在 chapter_a 之前），且导航页不是章节
    chapters = [n for n, is_ch in names if is_ch]
    assert chapters == ["chapter_b.xhtml", "chapter_a.xhtml"]
    assert ("nav.xhtml", False) in names


def test_docx_keeps_heading_body_and_table_in_document_order():
    result = normalize_sources([DOCX_FIXTURE])
    texts = [u.text for u in result.units]
    assert texts[0] == "第一场 茶楼"
    assert "甲走进茶楼，选了个靠窗的座位。" in texts
    assert any("玉牌" in t and "信物" in t for t in texts)
    # 表格行必须落在第二场之前（文档顺序，而非"先全文后全表"）
    table_idx = next(i for i, t in enumerate(texts) if "信物" in t)
    scene2_idx = texts.index("第二场 长街")
    assert table_idx < scene2_idx


def test_docx_ignores_header_and_footer_decoration():
    result = normalize_sources([DOCX_FIXTURE])
    assert "剧本草稿" not in result.markdown
    assert "页脚装饰" not in result.markdown
    assert "请勿外传" not in result.markdown


def test_docx_locators_use_heading_and_table_axes():
    result = normalize_sources([DOCX_FIXTURE])
    locators = {u.locator for u in result.units}
    assert "第一场 茶楼/标题" in locators
    assert "第一场 茶楼/段落1" in locators
    assert "第二场 长街/段落1" in locators
    assert "表格1/行1" in locators
    assert "表格1/行2" in locators


def test_docx_quote_traceable_into_normalized_markdown():
    result = normalize_sources([DOCX_FIXTURE])
    for unit in result.units:
        assert unit.text.strip() in result.markdown
    assert "乙递出一枚玉牌。" in result.markdown


def test_docx_manifest_reports_format_and_hash():
    import hashlib

    result = normalize_sources([DOCX_FIXTURE])
    meta = result.manifest["sources"][0]
    assert meta["format"] == "docx"
    assert meta["sha256"] == hashlib.sha256(DOCX_FIXTURE.read_bytes()).hexdigest()
    assert all(u.source_sha256 == meta["sha256"] for u in result.units)


def test_epub_follows_spine_order_not_filename():
    result = normalize_sources([EPUB_FIXTURE])
    texts = [u.text for u in result.units]
    assert "第二章 长街" in texts
    assert "第一章 茶楼" in texts
    # spine 里 chapter_b 在前，正文顺序必须照 spine（而不是 chapter_a 字母序在前）
    assert texts.index("第二章 长街") < texts.index("第一章 茶楼")


def test_epub_excludes_nav_and_scripts():
    result = normalize_sources([EPUB_FIXTURE])
    assert "目录" not in result.markdown
    assert "脚本内容不应进入正文" not in result.markdown
    assert "font-family" not in result.markdown


def test_epub_locators_use_chapter_file_and_paragraph():
    result = normalize_sources([EPUB_FIXTURE])
    locators = [u.locator for u in result.units]
    assert "chapter_b.xhtml#段落1" in locators
    assert "chapter_b.xhtml#段落2" in locators
    assert "chapter_a.xhtml#段落1" in locators
    assert not any("nav.xhtml" in loc for loc in locators)


def test_epub_quote_traceable_into_normalized_markdown():
    result = normalize_sources([EPUB_FIXTURE])
    assert "乙从屋檐上跃下。" in result.markdown
    for unit in result.units:
        assert unit.text.strip() in result.markdown


def test_multi_source_project_orders_docx_and_epub(tmp_path):
    order = tmp_path / "顺序清单.txt"
    result = normalize_sources([EPUB_FIXTURE, DOCX_FIXTURE])
    assert [Path(m["source_path"]).suffix for m in result.manifest["sources"]] == [".docx", ".epub"]
    assert result.manifest["sources"][0]["format"] == "docx"
    assert result.manifest["sources"][1]["format"] == "epub"
    _ = order


def test_normalization_dependencies_are_declared():
    """DOCX/EPUB 是正式输入，依赖必须写进 requirements.txt 与 pyproject.toml。"""
    req = (REPO_ROOT / "requirements.txt").read_text(encoding="utf-8")
    pyproj = (REPO_ROOT / "pyproject.toml").read_text(encoding="utf-8")
    for pkg in ("python-docx>=1.1,<2", "EbookLib>=0.18,<1", "beautifulsoup4>=4.12,<5"):
        assert pkg in req, f"requirements.txt 缺少 {pkg}"
        assert pkg in pyproj, f"pyproject.toml 缺少 {pkg}"
